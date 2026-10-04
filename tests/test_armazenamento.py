from datetime import datetime

import pytest
from openpyxl import load_workbook

from captador.armazenamento import (
    DadosInvalidos,
    HistoricoPlanilha,
    Usuario,
    neutralizar_formula,
)
from captador.clima import Leitura

ANA = Usuario("Ana", "ana@exemplo.com")


def test_cria_planilha_com_cabecalho_e_grava_registro(tmp_path):
    h = HistoricoPlanilha(tmp_path / "d.xlsx")
    h.salvar(Leitura(21.5, 63), ANA, agora=datetime(2026, 4, 8, 13, 5))

    linhas = list(load_workbook(h.caminho).active.iter_rows(values_only=True))
    assert linhas[0] == ("Data/Hora", "Temperatura (°C)", "Umidade (%)", "Nome", "E-mail")
    assert linhas[1] == ("08/04/2026 13:05", 21.5, 63, "Ana", "ana@exemplo.com")


def test_acrescenta_sem_apagar_o_historico(tmp_path):
    h = HistoricoPlanilha(tmp_path / "d.xlsx")
    for i in range(3):
        h.salvar(Leitura(20 + i, 50 + i), ANA, agora=datetime(2026, 4, 8, 10 + i, 0))

    registros = h.ultimos(10)
    assert [r.umidade for r in registros] == [50, 51, 52]
    assert registros[0].data_hora == "08/04/2026 10:00"


def test_ultimos_devolve_so_os_n_mais_recentes_em_ordem(tmp_path):
    h = HistoricoPlanilha(tmp_path / "d.xlsx")
    for i in range(15):
        h.salvar(Leitura(i, i), ANA, agora=datetime(2026, 4, 8, 8, i))

    registros = h.ultimos(10)
    assert len(registros) == 10
    assert registros[0].umidade == 5 and registros[-1].umidade == 14


def test_ultimos_sem_arquivo_devolve_lista_vazia(tmp_path):
    assert HistoricoPlanilha(tmp_path / "nao-existe.xlsx").ultimos() == []


def test_nao_deixa_arquivo_temporario_para_tras(tmp_path):
    h = HistoricoPlanilha(tmp_path / "d.xlsx")
    h.salvar(Leitura(20, 50), ANA)
    assert sorted(p.name for p in tmp_path.iterdir()) == ["d.xlsx"]


@pytest.mark.parametrize("malicioso", ['=HYPERLINK("http://x","clique")', "+cmd|' /C calc'!A0", "-2+3", "@SUM(A1)"])
def test_injecao_de_formula_e_neutralizada(tmp_path, malicioso):
    h = HistoricoPlanilha(tmp_path / "d.xlsx")
    h.salvar(Leitura(20, 50), Usuario(malicioso, "a@b.co"))

    celula = load_workbook(h.caminho).active["D2"]
    assert celula.data_type != "f", "não pode ser gravada como fórmula"
    assert celula.value == "'" + malicioso


def test_neutralizar_nao_mexe_em_texto_normal():
    assert neutralizar_formula("Maria da Silva") == "Maria da Silva"
    assert neutralizar_formula("") == ""


def test_validacao_do_usuario():
    assert Usuario.validar("  Ana  ", " ana@exemplo.com ") == ANA
    for nome, email in [
        ("", "a@b.co"),
        ("   ", "a@b.co"),
        ("Ana", ""),
        ("Ana", "sem-arroba"),
        ("Ana", "a@b"),
        ("Ana", "a b@c.co"),
        ("x" * 101, "a@b.co"),
        ("Ana", "a@" + "b" * 300 + ".co"),
    ]:
        with pytest.raises(DadosInvalidos):
            Usuario.validar(nome, email)
