"""Histórico em planilha Excel (.xlsx)."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .clima import Leitura

CABECALHO = ["Data/Hora", "Temperatura (°C)", "Umidade (%)", "Nome", "E-mail"]
FORMATO_DATA = "%d/%m/%Y %H:%M"
_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


class DadosInvalidos(ValueError):
    """Dado informado pelo usuário é inválido; a mensagem é segura para mostrar."""


@dataclass(frozen=True)
class Usuario:
    nome: str
    email: str

    @staticmethod
    def validar(nome: str, email: str) -> Usuario:
        nome, email = nome.strip(), email.strip()
        if not nome:
            raise DadosInvalidos("Informe o nome.")
        if len(nome) > 100:
            raise DadosInvalidos("O nome deve ter no máximo 100 caracteres.")
        if len(email) > 254 or not _EMAIL.match(email):
            raise DadosInvalidos("Informe um e-mail válido.")
        return Usuario(nome, email)


@dataclass(frozen=True)
class Registro:
    data_hora: str
    temperatura: float
    umidade: int
    nome: str
    email: str


def neutralizar_formula(texto: str) -> str:
    """Impede 'injeção de fórmula': o Excel executa células que começam com = + - @ (ex.: =HYPERLINK(...)).

    O prefixo ' faz o texto ser tratado como texto puro.
    """
    return "'" + texto if texto[:1] in ("=", "+", "-", "@", "\t", "\r") else texto


class HistoricoPlanilha:
    def __init__(self, caminho: Path) -> None:
        self.caminho = Path(caminho)

    def salvar(self, leitura: Leitura, usuario: Usuario, agora: datetime | None = None) -> None:
        agora = agora or datetime.now()
        if self.caminho.exists():
            planilha = load_workbook(self.caminho)
            aba = planilha.active
        else:
            planilha = Workbook()
            aba = planilha.active
            aba.append(CABECALHO)
        aba.append(
            [
                agora.strftime(FORMATO_DATA),
                leitura.temperatura,
                leitura.umidade,
                neutralizar_formula(usuario.nome),
                neutralizar_formula(usuario.email),
            ]
        )
        # Grava em arquivo temporário e troca: se o programa cair no meio, a planilha antiga não é corrompida.
        temporario = self.caminho.with_suffix(".tmp.xlsx")
        planilha.save(temporario)
        temporario.replace(self.caminho)

    def ultimos(self, quantidade: int = 10) -> list[Registro]:
        """Os últimos `quantidade` registros, do mais antigo para o mais novo."""
        if not self.caminho.exists():
            return []
        planilha = load_workbook(self.caminho, read_only=True)
        try:
            linhas = [linha for linha in planilha.active.iter_rows(min_row=2, values_only=True) if linha and linha[0]]
        finally:
            planilha.close()

        registros = []
        for linha in linhas[-quantidade:]:
            data, temperatura, umidade, nome, email = (tuple(linha) + (None,) * 5)[:5]
            registros.append(Registro(str(data), float(temperatura), int(umidade), str(nome or ""), str(email or "")))
        return registros
